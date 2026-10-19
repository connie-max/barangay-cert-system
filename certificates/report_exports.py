"""
Builds the downloadable Collections Report (Excel and PDF).

This file knows nothing about Django. The views hand it plain data:
    rows  = [{"date": date, "or_number": str, "payor": str,
              "nature": str, "amount": Decimal}, ...]
    start, end = datetime.date
so the same numbers shown on screen end up in both files.
"""
from decimal import Decimal
from io import BytesIO
from xml.sax.saxutils import escape

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.lib.enums import TA_CENTER, TA_RIGHT
from reportlab.platypus import (
    KeepTogether, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle,
)

# ---------------------------------------------------------------------------
# EDIT THESE: the header and signature lines printed on the report.
# (Change the text between the quotes, save, restart the server.)
# ---------------------------------------------------------------------------
HEADER_LINES = [
    "Republic of the Philippines",
    "Province of Davao de Oro",
    "Municipality of Nabunturan",
    "BARANGAY SASA",
]
# The name goes on the signature line; the title goes under it.
# Leave a name as "" to leave the line blank for hand-signing.
PREPARED_BY_NAME = ""
PREPARED_BY_TITLE = ""
NOTED_BY_NAME = ""
NOTED_BY_TITLE = ""

COLUMNS = ["Date", "OR No.", "Payor", "Nature of Collection", "Amount"]
BLUE = "2A5298"


def period_text(start, end):
    """'October 1, 2026 - October 31, 2026' (no %-d, so it works on Windows too)."""
    def fmt(d):
        return f"{d.strftime('%B')} {d.day}, {d.year}"
    return f"{fmt(start)} \u2013 {fmt(end)}"


def short_date(d):
    return f"{d.strftime('%b')} {d.day:02d}, {d.year}"


def money(value):
    return f"{Decimal(value):,.2f}"


def total_of(rows):
    return sum((Decimal(r["amount"]) for r in rows), Decimal("0.00"))


# ---------------------------------------------------------------------------
# Excel
# ---------------------------------------------------------------------------
def build_xlsx(rows, start, end):
    wb = Workbook()
    ws = wb.active
    ws.title = "Collections"

    thin = Side(style="thin", color="BBBBBB")
    box = Border(left=thin, right=thin, top=thin, bottom=thin)

    # Header block
    r = 1
    for line in HEADER_LINES:
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=5)
        c = ws.cell(row=r, column=1, value=line)
        c.alignment = Alignment(horizontal="center")
        c.font = Font(bold=(line == HEADER_LINES[-1]), size=11)
        r += 1
    r += 1
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=5)
    c = ws.cell(row=r, column=1, value="COLLECTIONS REPORT")
    c.font = Font(bold=True, size=14)
    c.alignment = Alignment(horizontal="center")
    r += 1
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=5)
    c = ws.cell(row=r, column=1, value=period_text(start, end))
    c.alignment = Alignment(horizontal="center")
    r += 2

    # Column headings
    head_row = r
    for col, name in enumerate(COLUMNS, start=1):
        c = ws.cell(row=r, column=col, value=name)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=BLUE)
        c.border = box
        c.alignment = Alignment(horizontal="right" if name == "Amount" else "left")
    r += 1

    # Data rows
    for row in rows:
        values = [row["date"], row["or_number"], row["payor"], row["nature"], float(row["amount"])]
        for col, v in enumerate(values, start=1):
            c = ws.cell(row=r, column=col, value=v)
            c.border = box
        ws.cell(row=r, column=1).number_format = "mmm dd, yyyy"
        ws.cell(row=r, column=1).alignment = Alignment(horizontal="left")
        ws.cell(row=r, column=5).number_format = '"\u20b1"#,##0.00'
        r += 1

    # Total row (a plain value, so it also shows in phone previews that don't calculate formulas)
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=4)
    ws.cell(row=r, column=1, value="TOTAL").font = Font(bold=True)
    tc = ws.cell(row=r, column=5, value=float(total_of(rows)))
    tc.font = Font(bold=True)
    tc.number_format = '"\u20b1"#,##0.00'
    for col in range(1, 6):
        ws.cell(row=r, column=col).border = box
        ws.cell(row=r, column=col).fill = PatternFill("solid", fgColor="F4F6FB")
    r += 3

    # Signature lines
    ws.cell(row=r, column=1, value="Prepared by:").font = Font(bold=True)
    ws.cell(row=r, column=4, value="Noted by:").font = Font(bold=True)
    r += 3
    ws.cell(row=r, column=1, value=PREPARED_BY_NAME or "______________________")
    ws.cell(row=r, column=4, value=NOTED_BY_NAME or "______________________")
    r += 1
    ws.cell(row=r, column=1, value=PREPARED_BY_TITLE)
    ws.cell(row=r, column=4, value=NOTED_BY_TITLE)

    # Widths, frozen headings, print setup
    for col, width in zip("ABCDE", (16, 18, 30, 30, 16)):
        ws.column_dimensions[col].width = width
    ws.freeze_panes = ws.cell(row=head_row + 1, column=1)
    ws.page_setup.orientation = "portrait"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = f"{head_row}:{head_row}"

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# PDF
# ---------------------------------------------------------------------------
def build_pdf(rows, start, end, generated_on):
    buf = BytesIO()
    doc = SimpleDocTemplate(
        buf, pagesize=A4,
        leftMargin=18 * mm, rightMargin=18 * mm, topMargin=16 * mm, bottomMargin=18 * mm,
        title="Collections Report", author="Barangay Certificate System",
    )

    center = ParagraphStyle("c", fontName="Helvetica", fontSize=10, leading=13, alignment=TA_CENTER)
    center_bold = ParagraphStyle("cb", parent=center, fontName="Helvetica-Bold", fontSize=12, leading=15)
    title = ParagraphStyle("t", parent=center, fontName="Helvetica-Bold", fontSize=14, leading=18, spaceBefore=10)
    cell = ParagraphStyle("cell", fontName="Helvetica", fontSize=9, leading=11)
    cell_right = ParagraphStyle("cr", parent=cell, alignment=TA_RIGHT)
    head = ParagraphStyle("h", parent=cell, fontName="Helvetica-Bold", textColor=colors.white)
    head_right = ParagraphStyle("hr", parent=head, alignment=TA_RIGHT)
    bold = ParagraphStyle("b", parent=cell, fontName="Helvetica-Bold")
    bold_right = ParagraphStyle("br", parent=bold, alignment=TA_RIGHT)
    small = ParagraphStyle("s", parent=cell, fontSize=9, leading=12)

    story = []
    for i, line in enumerate(HEADER_LINES):
        last = i == len(HEADER_LINES) - 1
        story.append(Paragraph(escape(line), center_bold if last else center))
    story.append(Paragraph("COLLECTIONS REPORT", title))
    story.append(Paragraph(escape(period_text(start, end)), center))
    story.append(Spacer(1, 8 * mm))

    # Table (Paragraph cells so long names wrap instead of running off the page)
    data = [[
        Paragraph("Date", head), Paragraph("OR No.", head), Paragraph("Payor", head),
        Paragraph("Nature of Collection", head), Paragraph("Amount (PHP)", head_right),
    ]]
    for row in rows:
        data.append([
            Paragraph(escape(short_date(row["date"])), cell),
            Paragraph(escape(row["or_number"]), cell),
            Paragraph(escape(row["payor"]), cell),
            Paragraph(escape(row["nature"]), cell),
            Paragraph(money(row["amount"]), cell_right),
        ])
    data.append([
        Paragraph("TOTAL", bold), "", "", "",
        Paragraph(money(total_of(rows)), bold_right),
    ])

    widths = [26 * mm, 30 * mm, 46 * mm, 44 * mm, 28 * mm]
    last_row = len(data) - 1
    table = Table(data, colWidths=widths, repeatRows=1)
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#" + BLUE)),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#BBBBBB")),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("TOPPADDING", (0, 0), (-1, -1), 5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
        ("BACKGROUND", (0, last_row), (-1, last_row), colors.HexColor("#F4F6FB")),
        ("SPAN", (0, last_row), (3, last_row)),
    ]))
    story.append(table)
    story.append(Spacer(1, 6 * mm))
    story.append(Paragraph(f"Total official receipts: {len(rows)}", small))

    # Signature block (kept together so it never splits across two pages)
    sig = Table(
        [
            [Paragraph("Prepared by:", bold), "", Paragraph("Noted by:", bold)],
            ["", "", ""],
            [Paragraph(escape(PREPARED_BY_NAME), center), "", Paragraph(escape(NOTED_BY_NAME), center)],
            [Paragraph(escape(PREPARED_BY_TITLE), center), "", Paragraph(escape(NOTED_BY_TITLE), center)],
        ],
        colWidths=[70 * mm, 34 * mm, 70 * mm],
        rowHeights=[6 * mm, 16 * mm, None, None],
    )
    sig.setStyle(TableStyle([
        ("LINEABOVE", (0, 2), (0, 2), 0.8, colors.black),
        ("LINEABOVE", (2, 2), (2, 2), 0.8, colors.black),
        ("TOPPADDING", (0, 2), (-1, 2), 3),
    ]))
    story.append(Spacer(1, 12 * mm))
    story.append(KeepTogether([sig]))

    def footer(canvas, doc_):
        canvas.saveState()
        canvas.setFont("Helvetica", 8)
        canvas.setFillColor(colors.HexColor("#777777"))
        canvas.drawString(18 * mm, 10 * mm, f"Generated {short_date(generated_on)}")
        canvas.drawRightString(A4[0] - 18 * mm, 10 * mm, f"Page {doc_.page}")
        canvas.restoreState()

    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    return buf.getvalue()